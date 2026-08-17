from unittest import TestCase
import io
import tempfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.services.document_conversion import PdfConverter
from app.services.document_conversion import DocxConverter
from app.services.document_conversion import DocumentReader
from tests.helpers.html import (
    generate_simple_html,
    generate_html_with_cyrillic,
    generate_html_with_table,
)


class TestPdfConverter(TestCase):
    @staticmethod
    def _convert_html(html: str) -> io.BytesIO:
        with tempfile.TemporaryDirectory() as temp_dir:
            html_path = Path(temp_dir) / "test.html"
            html_path.write_text(html, encoding="utf-8")

            converter = PdfConverter()

            return converter.convert(html_path)

    def test_convert_simple_html(self):
        html = generate_simple_html(
            title="Test document",
            text="Hello, world!",
        )

        result = TestPdfConverter._convert_html(html)

        reader = PdfReader(result)

        self.assertEqual(len(reader.pages), 1)

        text = reader.pages[0].extract_text()

        self.assertIn("Test document", text)
        self.assertIn("Hello, world!", text)

    def test_convert_cyrillic_html(self):
        html = generate_html_with_cyrillic()

        result = TestPdfConverter._convert_html(html)

        reader = PdfReader(result)

        text = reader.pages[0].extract_text()

        self.assertIn("Диссертационный совет", text)
        self.assertIn("Документ успешно сгенерирован.", text)

    def test_convert_html_with_table(self):
        html = generate_html_with_table()

        result = TestPdfConverter._convert_html(html)

        reader = PdfReader(result)

        text = reader.pages[0].extract_text()

        self.assertIn("First", text)
        self.assertIn("Second", text)
        self.assertIn("100", text)
        self.assertIn("200", text)


class TestDocxConverter(TestCase):
    @staticmethod
    def _convert_html(html: str) -> io.BytesIO:
        with tempfile.TemporaryDirectory() as temp_dir:
            html_path = Path(temp_dir) / "test.html"
            html_path.write_text(html, encoding="utf-8")

            converter = DocxConverter()

            return converter.convert(html_path)

    def test_convert_simple_html(self):
        html = generate_simple_html(
            title="Test document",
            text="Hello, world!",
        )

        result = TestDocxConverter._convert_html(html)

        self.assertIsNotNone(result)
        self.assertGreater(result.getbuffer().nbytes, 0)

        document = Document(result)

        text = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

        self.assertIn("Test document", text)
        self.assertIn("Hello, world!", text)

    def test_convert_cyrillic_html(self):
        html = generate_html_with_cyrillic()

        result = TestDocxConverter._convert_html(html)

        document = Document(result)

        text = "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
        )

        self.assertIn("Диссертационный совет", text)
        self.assertIn("Документ успешно сгенерирован.", text)

    def test_convert_html_with_table(self):
        html = generate_html_with_table()

        result = TestDocxConverter._convert_html(html)

        document = Document(result)

        table_text = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )

        self.assertIn("First", table_text)
        self.assertIn("Second", table_text)
        self.assertIn("100", table_text)
        self.assertIn("200", table_text)


class TestDocumentReader(TestCase):
    def test_read_html(self):
        html = generate_simple_html(
            title="Test document", text="This is a test")
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document.html"
            path.write_text(html, encoding="utf-8")

            result = DocumentReader.read_html(path)
            self.assertEqual(result, html)

    def test_read_html_without_html_extension(self):
        html = generate_simple_html(
            title="Test document", text="This is a test")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document"
            path.write_text(html, encoding="utf-8")

            result = DocumentReader.read_html(path)
            self.assertEqual(result, html)

    def test_reject_file_with_html_extension_but_invalid_content(self):
        garbage = """
        This is not an HTML document.
        Just some random text.
        """

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document.html"
            path.write_text(garbage, encoding="utf-8")

            with self.assertRaises(ValueError):
                DocumentReader.read_html(path)

    def test_file_not_found(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "does-not-exist.html"

            with self.assertRaises(FileNotFoundError):
                DocumentReader.read_html(path)

    def test_path_is_directory(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir)

            with self.assertRaises(ValueError):
                DocumentReader.read_html(path)

    def test_read_html_with_cyrillic(self):
        html = generate_html_with_cyrillic()

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document.html"
            path.write_text(html, encoding="utf-8")

            result = DocumentReader.read_html(path)

            self.assertIn("Диссертационный совет", result)
            self.assertIn("русский текст", result)